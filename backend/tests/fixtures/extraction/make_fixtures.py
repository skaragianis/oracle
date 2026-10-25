"""Regenerates the extraction fixtures. Not run by the tests: the fixtures are
committed bytes, so a library upgrade can't silently change what they contain.

    uv run python tests/fixtures/extraction/make_fixtures.py

Rerunning changes the bytes (random PDF /ID, docx timestamps) even though the
extracted text is identical, so rerun only when a fixture's content must change.

Regenerating goldens (ORACLE_UPDATE_GOLDENS=1) is a separate step.
"""

import re
import textwrap
from pathlib import Path

import docx
import pymupdf

OUT_DIR = Path(__file__).resolve().parent

PAGE_WIDTH, PAGE_HEIGHT = 612, 792
MARGIN = 72
FONT_SIZE = 10
LINE_HEIGHT = 13
PARAGRAPH_GAP = 9
WRAP_COLUMNS = 88

TOPICS = [
    ("lighthouse", "keeper", "oil lamp", "fog signal"),
    ("orchard", "grafting", "frost", "cider press"),
    ("telegraph", "relay station", "copper wire", "operator"),
    ("glacier", "moraine", "meltwater", "survey stake"),
    ("bakery", "sourdough", "proofing", "brick oven"),
    ("observatory", "refractor", "star chart", "dome shutter"),
    ("canal", "lock gate", "towpath", "barge horse"),
    ("tannery", "oak bark", "soaking pit", "hide"),
    ("monastery", "scriptorium", "vellum", "bell tower"),
    ("windmill", "millstone", "sail cloth", "grain hopper"),
    ("shipyard", "keel", "caulking", "slipway"),
    ("apiary", "brood frame", "smoker", "honey flow"),
]


SENTENCES_PER_PARAGRAPH = [4, 4, 2, 1]


def _paragraph(index: int) -> str:
    place, a, b, c = TOPICS[index % len(TOPICS)]
    round_number = index // len(TOPICS) + 1
    sentences = [
        f"Section {index + 1} of the survey describes the {place}, round "
        f"{round_number}.",
        f"Inspectors recorded the {a} in good order and noted that the {b} needed "
        f"attention before the next season.",
        f"The {c} was measured at {40 + index * 7} units and compared with the "
        f"figure from the previous visit, which had been {35 + index * 5} units.",
        f"Further details follow in the appendix for the {place} and its "
        f"neighbouring sites.",
    ]
    # Mixed lengths so the chunk overlap window spans a varying number of paragraphs
    return " ".join(sentences[: SENTENCES_PER_PARAGRAPH[index % 4]])


def _write_text_pdf(path: Path) -> None:
    doc = pymupdf.open()
    page = doc.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
    y = MARGIN
    for index in range(60):
        lines = textwrap.wrap(_paragraph(index), WRAP_COLUMNS)
        if index == 4:
            # A word broken across lines, as typesetters do.
            head, word = lines[0].rsplit(" ", 1)
            half = len(word) // 2
            lines[0] = f"{head} {word[:half]}-"
            lines[1] = f"{word[half:]} {lines[1]}"
        height = len(lines) * LINE_HEIGHT
        if y + height > PAGE_HEIGHT - MARGIN:
            page = doc.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
            y = MARGIN
        page.insert_text(
            (MARGIN, y), "\n".join(lines), fontsize=FONT_SIZE, lineheight=1.3
        )
        y += height + PARAGRAPH_GAP
    doc.save(path, deflate=True)
    doc.close()


def _write_malformed_pdf(path: Path) -> None:
    doc = pymupdf.open()
    for number in (1, 2, 3):
        page = doc.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
        page.insert_text(
            (MARGIN, MARGIN),
            f"Malformed fixture page {number}.\nThe cross reference table of this "
            f"file is broken\nso MuPDF has to rebuild it by scanning the objects.",
            fontsize=FONT_SIZE,
        )
    # An unterminated array in page 2's content stream, as some producers emit.
    xref = doc[1].get_contents()[0]
    doc.update_stream(xref, doc.xref_stream(xref) + b"\nQQ bogus [ /Nope")
    data = doc.tobytes(deflate=False, garbage=0)
    doc.close()

    # Point startxref at the wrong offset and drop the xref keyword, forcing a
    # repair pass at open time.
    data = re.sub(rb"startxref\s+\d+", b"startxref\n7", data)
    data = data.replace(b"\nxref\n", b"\nxrex\n", 1)
    path.write_bytes(data)


def _write_scanned_pdf(path: Path) -> None:
    source = pymupdf.open()
    page = source.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
    page.insert_text(
        (MARGIN, MARGIN),
        "This page is only a picture of text.\nThere is no text layer to extract.",
        fontsize=18,
    )
    pixmap = page.get_pixmap(dpi=72, colorspace=pymupdf.csGRAY)
    png = pixmap.tobytes("png")
    source.close()

    doc = pymupdf.open()
    page = doc.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
    page.insert_image(page.rect, stream=png)
    doc.save(path, deflate=True)
    doc.close()


def _write_tables_pdf(path: Path) -> None:
    doc = pymupdf.open()
    page = doc.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
    page.insert_text((MARGIN, MARGIN), "Quarterly stock summary", fontsize=14)
    header = ["Item", "Region", "Units", "Price"]
    rows = [
        ["Lantern", "North", "120", "4.50"],
        ["Rope", "South", "75", "2.25"],
        ["Anchor", "East", "18", "61.00"],
        ["Tarpaulin", "West", "42", "12.80"],
    ]
    cell_width, cell_height = 110, 24
    top = MARGIN + 30
    for row_index, row in enumerate([header, *rows]):
        for col_index, text in enumerate(row):
            rect = pymupdf.Rect(
                MARGIN + col_index * cell_width,
                top + row_index * cell_height,
                MARGIN + (col_index + 1) * cell_width,
                top + (row_index + 1) * cell_height,
            )
            page.draw_rect(rect, width=0.8)
            page.insert_text(
                (rect.x0 + 6, rect.y1 - 8),
                text,
                fontsize=FONT_SIZE,
                fontname="hebo" if row_index == 0 else "helv",
            )
    page.insert_text(
        (MARGIN, top + 6 * cell_height),
        "Totals are reconciled at the end of each quarter.",
        fontsize=FONT_SIZE,
    )
    doc.save(path, deflate=True)
    doc.close()


def _write_nested_table_docx(path: Path) -> None:
    document = docx.Document()
    document.add_paragraph("Equipment register introduction paragraph.")
    outer = document.add_table(rows=2, cols=2)
    outer.cell(0, 0).text = "Outer header left"
    outer.cell(0, 1).text = "Outer header right"
    outer.cell(1, 0).text = "Outer body left"
    holder = outer.cell(1, 1)
    holder.paragraphs[0].text = "Cell text before the nested table"
    inner = holder.add_table(rows=2, cols=2)
    inner.cell(0, 0).text = "Inner alpha"
    inner.cell(0, 1).text = "Inner beta"
    inner.cell(1, 0).text = "Inner gamma"
    inner.cell(1, 1).text = "Inner delta"
    holder.add_paragraph("Cell text after the nested table")
    document.add_paragraph("Closing paragraph after the tables.")
    document.save(str(path))


def main() -> None:
    _write_text_pdf(OUT_DIR / "text.pdf")
    _write_malformed_pdf(OUT_DIR / "malformed.pdf")
    _write_scanned_pdf(OUT_DIR / "scanned.pdf")
    _write_tables_pdf(OUT_DIR / "tables.pdf")
    _write_nested_table_docx(OUT_DIR / "nested_table.docx")


if __name__ == "__main__":
    main()
